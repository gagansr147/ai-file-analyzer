import streamlit as st
from google import genai
from google.genai import types
from docx import Document
from pypdf import PdfReader
from io import BytesIO
import html

# --------------------------------------------------
# PAGE CONFIGURATION
# --------------------------------------------------
st.set_page_config(
    page_title="DocuMind AI | Document Intelligence",
    page_icon="✦",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --------------------------------------------------
# CUSTOM CSS - PREMIUM PURPLE & WHITE UI
# --------------------------------------------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap');

:root {
    --primary: #7047EB;
    --primary-dark: #5632CE;
    --ink: #211B38;
    --muted: #79758B;
    --border: #EAE6F4;
    --surface: #F8F7FC;
}

html, body, [class*="css"] {
    font-family: 'DM Sans', sans-serif;
}

.stApp {
    background: #FBFAFE;
    color: #211B38;
}

[data-testid="stHeader"] {
    background: rgba(251,250,254,0.92);
}

[data-testid="stMainBlockContainer"] {
    max-width: 1320px;
    padding-top: 1.8rem;
    padding-bottom: 4rem;
}

h1, h2, h3 {
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    color: #211B38 !important;
    letter-spacing: -0.7px;
}

p, label, .stCaption {
    color: #79758B;
}

[data-testid="stSidebar"] {
    background: #F3F0FC;
    border-right: 1px solid #E8E1F5;
}

[data-testid="stSidebar"] > div:first-child {
    padding-top: 1.5rem;
}

.brand {
    display: flex;
    align-items: center;
    gap: 12px;
    margin-bottom: 32px;
}

.brand-icon {
    width: 46px;
    height: 46px;
    border-radius: 15px;
    background: linear-gradient(135deg,#8D6BFF,#613CE1);
    color: white;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 25px;
    box-shadow: 0 7px 18px #7047EB35;
}

.brand-name {
    font-family: 'Plus Jakarta Sans',sans-serif;
    font-size: 19px;
    font-weight: 800;
    color: #261B48;
}

.brand-caption {
    font-size: 11px;
    color: #89829F;
    margin-top: 2px;
}

.sidebar-label {
    color: #9891AA;
    font-size: 10px;
    font-weight: 700;
    letter-spacing: 1.7px;
    margin: 24px 0 9px;
}

.sidebar-note {
    background: white;
    border: 1px solid #E9E2F7;
    border-radius: 15px;
    padding: 15px;
    margin-top: 28px;
    font-size: 12px;
    color: #77718A;
    line-height: 1.7;
}

.hero {
    padding: 8px 0 22px;
}

.eyebrow {
    display: inline-block;
    background: #F0EAFF;
    border: 1px solid #E3D8FF;
    color: #6943D8;
    padding: 7px 12px;
    border-radius: 30px;
    font-size: 11px;
    font-weight: 700;
    letter-spacing: .5px;
    margin-bottom: 17px;
}

.hero-title {
    font-family: 'Plus Jakarta Sans',sans-serif;
    font-size: clamp(30px, 4vw, 43px);
    line-height: 1.2;
    letter-spacing: -1.7px;
    font-weight: 800;
    color: #241B3D;
    margin: 0 0 12px;
}

.hero-title span {
    color: #7047EB;
}

.hero-subtitle {
    color: #827C94;
    font-size: 15px;
    line-height: 1.8;
    max-width: 620px;
}

.status-pill {
    display: inline-flex;
    align-items: center;
    gap: 8px;
    background: white;
    border: 1px solid #E9E5F2;
    padding: 10px 13px;
    border-radius: 30px;
    color: #625C74;
    font-size: 12px;
    white-space: nowrap;
    margin-top: 8px;
}

.status-dot {
    width: 8px;
    height: 8px;
    background: #28B77B;
    border-radius: 50%;
}

.upload-card {
    background: white;
    border: 1px solid #EAE4F5;
    border-radius: 22px;
    padding: 25px 25px 20px;
    box-shadow: 0 8px 32px #35215B08;
    margin: 12px 0 22px;
}

.section-kicker {
    color: #8D84A3;
    font-size: 10px;
    font-weight: 700;
    letter-spacing: 1.6px;
    text-transform: uppercase;
    margin-bottom: 8px;
}

.section-title {
    color: #29213E;
    font-family: 'Plus Jakarta Sans',sans-serif;
    font-size: 21px;
    font-weight: 750;
    margin-bottom: 6px;
}

.section-description {
    font-size: 13px;
    color: #8A849A;
    line-height: 1.7;
}

.upload-icon {
    width: 62px;
    height: 62px;
    border-radius: 19px;
    background: linear-gradient(135deg,#F0E9FF,#F8F5FF);
    border: 1px solid #E7DEFF;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 30px;
    margin: 3px auto 13px;
}

.feature-card {
    background: white;
    border: 1px solid #EBE6F5;
    border-radius: 18px;
    padding: 19px;
    min-height: 150px;
    box-shadow: 0 5px 20px #35215B05;
}

.feature-icon {
    display: flex;
    align-items: center;
    justify-content: center;
    width: 41px;
    height: 41px;
    border-radius: 13px;
    background: #F1ECFF;
    font-size: 21px;
    margin-bottom: 14px;
}

.feature-title {
    color: #2A2240;
    font-weight: 700;
    font-size: 14px;
    margin-bottom: 7px;
}

.feature-description {
    color: #89839A;
    font-size: 12px;
    line-height: 1.7;
}

.panel {
    background: white;
    border: 1px solid #EAE5F4;
    border-radius: 19px;
    padding: 23px;
    box-shadow: 0 6px 25px #35215B06;
    margin: 12px 0 20px;
}

.file-chip {
    display: flex;
    align-items: center;
    gap: 12px;
    background: #F8F6FE;
    border: 1px solid #EDE6FA;
    border-radius: 13px;
    padding: 13px;
    margin: 12px 0;
}

.file-symbol {
    display: flex;
    align-items: center;
    justify-content: center;
    width: 43px;
    height: 43px;
    flex-shrink: 0;
    border-radius: 12px;
    background: #EDE5FF;
    font-size: 22px;
}

.file-name {
    color: #34284F;
    font-size: 13px;
    font-weight: 700;
    overflow-wrap: anywhere;
}

.file-meta {
    color: #9189A3;
    font-size: 11px;
    margin-top: 4px;
}

.result-card {
    background: white;
    border: 1px solid #E9E3F4;
    border-radius: 19px;
    padding: 23px;
    margin: 12px 0 22px;
}

.result-label {
    color: #8A81A0;
    font-size: 10px;
    font-weight: 700;
    letter-spacing: 1.5px;
    margin-bottom: 7px;
}

.chat-intro {
    background: linear-gradient(120deg,#F3EEFF,#FBF9FF);
    border: 1px solid #E9E0FF;
    border-radius: 17px;
    padding: 17px;
    margin: 12px 0 18px;
    color: #65577F;
    font-size: 13px;
    line-height: 1.8;
}

div.stButton > button {
    border-radius: 12px;
    min-height: 43px;
    font-weight: 650;
    transition: all .2s ease;
}

div.stButton > button[kind="primary"] {
    background: linear-gradient(110deg,#7850F0,#623BE0);
    border: none;
    color: white;
    box-shadow: 0 5px 14px #7047EB25;
}

div.stButton > button[kind="primary"]:hover {
    background: linear-gradient(110deg,#6840E0,#5129C7);
    box-shadow: 0 8px 20px #7047EB35;
    transform: translateY(-1px);
}

/* --------------------------------------------------
   CHAT MESSAGES: FORCE DARK TEXT
   -------------------------------------------------- */

div[data-testid="stChatMessage"] {
    background: #FFFFFF !important;
    color: #211B38 !important;
    border: 1px solid #EEE9F6;
    border-radius: 16px;
    padding: 15px;
    margin-bottom: 12px;
}

/* Message content, paragraphs, lists and inline formatting */
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"],
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] p,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] li,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] ol,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] ul,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] span,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] strong,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] em {
    color: #211B38 !important;
    -webkit-text-fill-color: #211B38 !important;
}

/* Headings in AI responses */
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] h1,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] h2,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] h3,
div[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] h4 {
    color: #35245E !important;
    -webkit-text-fill-color: #35245E !important;
}

/* Code in AI responses */
div[data-testid="stChatMessage"] pre,
div[data-testid="stChatMessage"] code {
    color: #35245E !important;
}

/* --------------------------------------------------
   CHAT INPUT: PURPLE TEXT + WHITE BACKGROUND
   -------------------------------------------------- */

[data-testid="stChatInput"] {
    background: #FFFFFF !important;
    border: 1px solid #DCD0FA !important;
    border-radius: 16px !important;
    box-shadow: 0 4px 16px rgba(112, 71, 235, 0.08);
}

[data-testid="stChatInput"]:focus-within {
    border-color: #7047EB !important;
    box-shadow: 0 0 0 3px rgba(112, 71, 235, 0.12) !important;
}

[data-testid="stChatInput"] textarea,
[data-testid="stChatInput"] textarea:focus {
    color: #5632CE !important;
    -webkit-text-fill-color: #5632CE !important;
    caret-color: #7047EB !important;
    background-color: #FFFFFF !important;
    font-family: 'DM Sans', sans-serif !important;
    font-size: 14px !important;
}

[data-testid="stChatInput"] textarea::placeholder {
    color: #968BAF !important;
    -webkit-text-fill-color: #968BAF !important;
    opacity: 1 !important;
}

[data-testid="stChatInput"] button {
    color: #7047EB !important;
    background: #F1ECFF !important;
    border-radius: 10px !important;
}

[data-testid="stChatInput"] button:hover {
    color: #FFFFFF !important;
    background: #7047EB !important;
}

/* File uploader */
[data-testid="stFileUploader"] {
    background: #FCFBFF;
    border: 1px dashed #CDBDFA;
    border-radius: 16px;
    padding: 15px;
}

[data-testid="stFileUploader"] section {
    background: transparent;
}

.stTabs [data-baseweb="tab-list"] {
    gap: 10px;
    border-bottom: 1px solid #EAE5F3;
}

.stTabs [data-baseweb="tab"] {
    background: transparent;
    border-radius: 10px 10px 0 0;
    padding: 10px 15px;
}

hr {
    border-color: #EEE9F6;
}

div[data-testid="stAlert"] {
    border-radius: 13px;
}

.footer {
    color: #A39CB2;
    font-size: 11px;
    text-align: center;
    padding: 28px 0 5px;
}

@media (max-width: 700px) {
    [data-testid="stMainBlockContainer"] {
        padding: 1rem 1rem 3rem;
    }

    .hero-title {
        font-size: 30px;
    }

    .upload-card, .panel, .result-card {
        padding: 17px;
    }
}
</style>
""", unsafe_allow_html=True)


# --------------------------------------------------
# GEMINI CLIENT
# --------------------------------------------------
try:
    api_key = st.secrets["GEMINI_API_KEY"]
    client = genai.Client(api_key=api_key)
except Exception:
    st.error(
        "Gemini could not be initialized. Check your "
        "GEMINI_API_KEY in .streamlit/secrets.toml."
    )
    st.stop()


# --------------------------------------------------
# SESSION STATE
# --------------------------------------------------
defaults = {
    "file_id": None,
    "file_content": None,
    "file_name": None,
    "messages": [],
    "summary": None,
    "summary_error": None,
    "chat_error": None,
    "suggested_question": None,
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# --------------------------------------------------
# DOCUMENT EXTRACTION
# --------------------------------------------------
def extract_content(uploaded_file):
    filename = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )
        return {"text": text, "image": None}

    if filename.endswith(".docx"):
        doc = Document(BytesIO(data))
        paragraphs = [
            p.text for p in doc.paragraphs if p.text.strip()
        ]

        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return {"text": "\n".join(paragraphs), "image": None}

    if filename.endswith(".txt"):
        return {
            "text": data.decode("utf-8", errors="replace"),
            "image": None,
        }

    return {
        "text": "",
        "image": types.Part.from_bytes(
            data=data,
            mime_type=uploaded_file.type or "image/png",
        ),
    }


# --------------------------------------------------
# GEMINI REQUEST
# --------------------------------------------------
def ask_ai(file_content, question, history):
    document_text = file_content["text"]

    if not document_text.strip() and file_content["image"] is None:
        raise ValueError(
            "No readable text was found in this file. "
            "If this is a scanned PDF, OCR may be required."
        )

    prompt = f"""
You are DocuMind AI, a helpful document assistant.

Use the uploaded file as your primary source.
Answer in clear, accurate, easy-to-understand English.
Use headings, lists and examples when they help.
Never invent information that is not in the document.
If the answer is missing from the document, say so.
For interview questions, create relevant questions based on the file.

DOCUMENT CONTENT:
{document_text[:60000]}

CURRENT QUESTION:
{question}
"""

    contents = []

    for message in history[-8:]:
        contents.append(
            types.Content(
                role=message["role"],
                parts=[
                    types.Part.from_text(text=message["content"])
                ],
            )
        )

    parts = [types.Part.from_text(text=prompt)]

    if file_content["image"] is not None:
        parts.append(file_content["image"])

    contents.append(types.Content(role="user", parts=parts))

    response = client.models.generate_content(
        model="gemini-3.8-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            temperature=0.3,
        ),
    )

    if not response.text:
        raise RuntimeError(
            "Gemini returned an empty response. Please try again."
        )

    return response.text


# --------------------------------------------------
# SIDEBAR
# --------------------------------------------------
with st.sidebar:
    st.markdown("""
    <div class="brand">
        <div class="brand-icon">✦</div>
        <div>
            <div class="brand-name">DocuMind AI</div>
            <div class="brand-caption">DOCUMENT INTELLIGENCE</div>
        </div>
    </div>
    <div class="sidebar-label">YOUR WORKSPACE</div>
    """, unsafe_allow_html=True)

    st.markdown("**✧  Document analyzer**")
    st.caption("Upload, summarize and explore your files.")

    st.markdown(
        '<div class="sidebar-label">QUICK GUIDE</div>',
        unsafe_allow_html=True
    )

    st.markdown("""
    <div style="font-size:13px;line-height:2.1;color:#6F6883">
    <b>01</b> &nbsp; Upload a document<br>
    <b>02</b> &nbsp; Generate its summary<br>
    <b>03</b> &nbsp; Ask questions about it
    </div>
    <div class="sidebar-note">
        <b style="color:#47366F">✦ Powered by Gemini</b><br>
        Your AI study and document companion for extracting
        key insights, understanding concepts and preparing
        interview questions.
    </div>
    """, unsafe_allow_html=True)

    if st.button("＋ Start a new analysis", use_container_width=True):
        for key, value in defaults.items():
            st.session_state[key] = value
        st.rerun()

    st.markdown("""
    <div style="font-size:10px;color:#A19AB2;margin-top:28px">
    DOCUMIND AI · PERSONAL WORKSPACE
    </div>
    """, unsafe_allow_html=True)


# --------------------------------------------------
# HERO HEADER
# --------------------------------------------------
st.markdown("""
<div class="hero">
    <div class="eyebrow">✦ YOUR AI-POWERED WORKSPACE</div>
    <div class="hero-title">
        Your documents,<br><span>beautifully understood.</span>
    </div>
    <div class="hero-subtitle">
        Turn complex documents into clear insights.
        Generate smart summaries, explore important details,
        and get answers directly from your files.
    </div>
</div>
""", unsafe_allow_html=True)

st.markdown("""
<div style="display:flex;justify-content:flex-end;margin-top:-5px">
    <div class="status-pill">
        <span class="status-dot"></span> AI assistant ready
    </div>
</div>
""", unsafe_allow_html=True)


# --------------------------------------------------
# UPLOAD SECTION
# --------------------------------------------------
st.markdown("""
<div class="upload-card">
    <div class="section-kicker">GET STARTED</div>
    <div class="section-title">Bring your document to life</div>
    <div class="section-description">
        Upload a file to unlock its summary, key information
        and an interactive AI conversation.
    </div>
    <div class="upload-icon">⇧</div>
    <div style="text-align:center;font-size:14px;
                color:#43365D;font-weight:700">
        Drop your file below or browse your computer
    </div>
    <div style="text-align:center;font-size:12px;
                color:#9690A6;margin:7px 0 5px">
        PDF · DOCX · TXT · PNG · JPG · WEBP
    </div>
</div>
""", unsafe_allow_html=True)

uploaded_file = st.file_uploader(
    "Choose a file",
    type=["pdf", "docx", "txt", "png", "jpg", "jpeg", "webp"],
    label_visibility="collapsed",
    help="Choose a PDF, Word document, text file or image.",
)


# --------------------------------------------------
# EMPTY WORKSPACE
# --------------------------------------------------
if uploaded_file is None:
    st.markdown("""
    <div style="height:5px"></div>
    <div class="section-kicker">WHAT YOU CAN DO</div>
    """, unsafe_allow_html=True)

    col1, col2 = st.columns(2)

    with col1:
        st.markdown("""
        <div class="feature-card">
            <div class="feature-icon">✧</div>
            <div class="feature-title">Smart summaries</div>
            <div class="feature-description">
                Understand the main ideas, key points,
                important details and conclusions without
                reading every line.
            </div>
        </div>
        """, unsafe_allow_html=True)

    with col2:
        st.markdown("""
        <div class="feature-card">
            <div class="feature-icon">☏</div>
            <div class="feature-title">Chat with your file</div>
            <div class="feature-description">
                Ask questions, simplify technical concepts,
                discover important facts and prepare
                interview questions from your document.
            </div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("""
    <div class="footer">
        Made for clearer thinking · Powered by Gemini
    </div>
    """, unsafe_allow_html=True)


# --------------------------------------------------
# FILE PROCESSING
# --------------------------------------------------
else:
    file_id = f"{uploaded_file.name}:{uploaded_file.size}"

    if st.session_state.file_id != file_id:
        try:
            extracted = extract_content(uploaded_file)

            st.session_state.file_id = file_id
            st.session_state.file_content = extracted
            st.session_state.file_name = uploaded_file.name
            st.session_state.messages = []
            st.session_state.summary = None
            st.session_state.summary_error = None
            st.session_state.chat_error = None
            st.session_state.suggested_question = None

        except Exception as error:
            st.error("We couldn't read this file.")
            st.code(f"{type(error).__name__}: {error}")
            st.stop()

    file_content = st.session_state.file_content

    safe_name = html.escape(uploaded_file.name)
    size_kb = uploaded_file.size / 1024

    st.markdown(f"""
    <div class="panel">
        <div class="section-kicker">CURRENT DOCUMENT</div>
        <div class="file-chip">
            <div class="file-symbol">▤</div>
            <div style="flex:1;min-width:0">
                <div class="file-name">{safe_name}</div>
                <div class="file-meta">
                    {size_kb:.1f} KB &nbsp; · &nbsp; Uploaded successfully
                </div>
            </div>
            <div style="color:#24966B;font-size:12px;font-weight:700">
                ✓ Ready
            </div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    if file_content["image"] is not None:
        with st.expander("Preview uploaded image"):
            st.image(uploaded_file, use_container_width=True)

    with st.expander("View extracted document text"):
        if file_content["text"].strip():
            st.text(file_content["text"][:20000])
        else:
            st.info(
                "No text was extracted. The uploaded image will "
                "be sent to Gemini for analysis."
            )

    # --------------------------------------------------
    # SUMMARY
    # --------------------------------------------------
    st.markdown("""
    <div class="section-kicker">01 / DOCUMENT INSIGHTS</div>
    <div class="section-title">Your smart summary</div>
    <div class="section-description">
        Extract the essential information from your document.
    </div>
    """, unsafe_allow_html=True)

    if st.button(
        "✧  Generate AI summary",
        type="primary",
        use_container_width=True,
    ):
        st.session_state.summary_error = None

        try:
            with st.spinner(
                "Reading your document and preparing insights..."
            ):
                summary = ask_ai(
                    file_content,
                    (
                        "Summarize this document. Explain its purpose, "
                        "main topics, key points, important details, "
                        "and conclusions. Use clear headings and bullets. "
                        "Keep the summary concise but useful."
                    ),
                    [],
                )
                st.session_state.summary = summary

        except Exception as error:
            st.session_state.summary_error = (
                f"{type(error).__name__}: {error}"
            )

    if st.session_state.summary_error:
        st.error("Summary generation failed.")
        st.code(st.session_state.summary_error)

    if st.session_state.summary:
        st.markdown("""
        <div class="result-card">
            <div class="result-label">✧ AI-GENERATED DOCUMENT SUMMARY</div>
        </div>
        """, unsafe_allow_html=True)

        # Force summary text to remain readable
        st.markdown("""
        <style>
        [data-testid="stMainBlockContainer"] .stMarkdown,
        [data-testid="stMainBlockContainer"] .stMarkdown p,
        [data-testid="stMainBlockContainer"] .stMarkdown li {
            color: #211B38;
        }
        </style>
        """, unsafe_allow_html=True)

        st.markdown(st.session_state.summary)

        st.download_button(
            "⬇ Download summary as TXT",
            data=st.session_state.summary,
            file_name="document_summary.txt",
            mime="text/plain",
            use_container_width=True,
        )

    # --------------------------------------------------
    # DOCUMENT CHAT
    # --------------------------------------------------
    st.divider()

    st.markdown("""
    <div class="section-kicker">02 / ASK YOUR DOCUMENT</div>
    <div class="section-title">Let's explore your file</div>
    <div class="section-description">
        Ask a question below. Follow-up questions use your
        recent conversation for context.
    </div>
    <div class="chat-intro">
        <b>✦ Your AI assistant is ready.</b><br>
        Try asking for a simple explanation, the key concepts,
        or five interview questions based on your document.
    </div>
    """, unsafe_allow_html=True)

    st.caption("QUICK QUESTIONS")

    q1, q2, q3 = st.columns(3)

    with q1:
        if st.button("Explain simply", use_container_width=True):
            st.session_state.suggested_question = (
                "Explain the main concepts in this document "
                "in simple English."
            )

    with q2:
        if st.button("Key takeaways", use_container_width=True):
            st.session_state.suggested_question = (
                "What are the five most important takeaways "
                "from this document?"
            )

    with q3:
        if st.button("Interview questions", use_container_width=True):
            st.session_state.suggested_question = (
                "Create five useful interview questions "
                "and answers based on this document."
            )

    # Show conversation history
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    question = st.chat_input(
        "Ask anything about your document..."
    )

    if question is None and st.session_state.suggested_question:
        question = st.session_state.suggested_question
        st.session_state.suggested_question = None

    if question:
        st.session_state.chat_error = None

        st.session_state.messages.append({
            "role": "user",
            "content": question,
        })

        with st.chat_message("user"):
            st.markdown(question)

        try:
            with st.chat_message("assistant"):
                with st.spinner("Thinking through your question..."):
                    answer = ask_ai(
                        file_content,
                        question,
                        st.session_state.messages[:-1],
                    )
                    st.markdown(answer)

            st.session_state.messages.append({
                "role": "assistant",
                "content": answer,
            })

        except Exception as error:
            st.session_state.messages.pop()

            st.session_state.chat_error = (
                f"{type(error).__name__}: {error}"
            )

        if st.session_state.chat_error:
            st.error("The AI couldn't answer this question.")
            st.code(st.session_state.chat_error)

    st.markdown("""
    <div class="footer">
        DocuMind AI · Your documents, understood.
    </div>
    """, unsafe_allow_html=True)
